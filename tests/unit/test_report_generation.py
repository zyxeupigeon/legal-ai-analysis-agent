import json
import re

import pytest
from docx import Document
from docx.oxml.ns import qn

from services.api.app.models import (
    EvidenceSource,
    IdentityGate,
    RelatedProduct,
    ReportDraft,
    ReportInsightDraft,
    ReportSectionDraft,
    ReportSourceRecord,
    Task,
    TaskStatus,
)
from services.api.app.provider_errors import ProviderInvalidResponse
from services.api.app.report_document import build_report_docx
from services.api.app.report_provider import (
    REPORT_SECTION_HEADINGS,
    EvidenceBoundReportDraftProvider,
    FixtureReportDraftProvider,
)
from services.api.app.search_provider import SearchDocument


class FakeModel:
    ready = True
    model_name = "fixture-report-model"
    provider_name = "fixture"

    def __init__(self, payload: dict) -> None:
        self.payload = payload

    def generate_json(self, system_prompt: str, user_prompt: str) -> str:
        if "证据相关性审计" in user_prompt:
            match = re.search(r"待审目标：\n(.*?)\n\n任务卡片与实际公开证据", user_prompt, re.S)
            assert match
            targets = json.loads(match.group(1))
            return json.dumps(
                {
                    "items": [
                        {
                            "target_id": item["target_id"],
                            "status": "supported",
                            "supported_urls": item["candidate_urls"],
                            "reason": "虚构测试中视为直接支持。",
                        }
                        for item in targets
                    ]
                },
                ensure_ascii=False,
            )
        assert "不得使用模型记忆" in system_prompt
        assert "actual.example/product" in user_prompt
        return json.dumps(self.payload, ensure_ascii=False)


class SequenceModel(FakeModel):
    def __init__(self, payloads: list[dict], audit_statuses: list[str] | None = None) -> None:
        self.payloads = list(payloads)
        self.audit_statuses = list(audit_statuses or [])
        self.prompts: list[str] = []

    def generate_json(self, system_prompt: str, user_prompt: str) -> str:
        self.prompts.append(user_prompt)
        if "证据相关性审计" in user_prompt:
            if self.audit_statuses:
                match = re.search(r"待审目标：\n(.*?)\n\n任务卡片与实际公开证据", user_prompt, re.S)
                assert match
                targets = json.loads(match.group(1))
                status = self.audit_statuses.pop(0)
                return json.dumps(
                    {
                        "items": [
                            {
                                "target_id": item["target_id"],
                                "status": status,
                                "supported_urls": item["candidate_urls"] if status == "supported" else [],
                                "reason": "测试证据相关性结果。",
                            }
                            for item in targets
                        ]
                    },
                    ensure_ascii=False,
                )
            return FakeModel.generate_json(self, system_prompt, user_prompt)
        return json.dumps(self.payloads.pop(0), ensure_ascii=False)


class FakeSearch:
    ready = True
    provider_name = "fixture-search"
    search_tier = "test_only"

    def search(self, query: str, *, max_results: int = 5) -> list[SearchDocument]:
        return [
            SearchDocument(
                query=query,
                title="产品官方说明",
                url="https://actual.example/product",
                content="该文本仅用于测试证据约束。",
                retrieval_status="readable",
            )
        ]


def fixture_task() -> Task:
    from datetime import UTC, datetime

    now = datetime.now(UTC)
    return Task(
        task_id="fixture-task",
        submitted_product_name="衡法智研",
        canonical_product_name="衡法智研",
        operator_name="衡法科技",
        product_type="法律AI研究产品",
        task_status=TaskStatus.RESEARCHING,
        identity_gate=IdentityGate.CONFIRMED,
        status_message="测试",
        created_at=now,
        updated_at=now,
    )


def report_payload(source_urls: list[str]) -> dict:
    return {
        "report_title": "会被系统覆盖",
        "report_subtitle": "会被系统覆盖",
        "executive_summary": ["测试摘要。"],
        "executive_summary_source_urls": source_urls,
        "sections": [
            {
                "heading": f"模型自拟标题{index}",
                "paragraphs": ["只描述测试证据。"],
                "key_points": ["测试要点。"],
                "source_urls": source_urls,
            }
            for index in range(5)
        ],
        "insights": [
            {
                "title": "测试洞察",
                "analysis": "只验证引用过滤。",
                "source_urls": source_urls,
            }
        ],
        "limitations": [],
        "sources": [],
    }


def test_report_provider_keeps_only_retrieved_sources_and_fixed_headings() -> None:
    provider = EvidenceBoundReportDraftProvider(
        FakeModel(report_payload(["https://actual.example/product", "https://invented.example"])),
        FakeSearch(),
    )
    draft = provider.generate(fixture_task())

    assert [section.heading for section in draft.sections] == REPORT_SECTION_HEADINGS
    assert draft.sections[0].source_urls == ["https://actual.example/product"]
    assert [source.url for source in draft.sources] == ["https://actual.example/product"]
    assert any("免费测试级公开检索" in item for item in draft.limitations)


def test_report_provider_rejects_draft_without_real_citation() -> None:
    provider = EvidenceBoundReportDraftProvider(
        FakeModel(report_payload(["https://invented.example"])),
        FakeSearch(),
    )
    with pytest.raises(ProviderInvalidResponse, match="未完成主张—证据绑定"):
        provider.generate(fixture_task())


def test_report_provider_repairs_missing_claim_citations_once() -> None:
    missing = report_payload([])
    valid = report_payload(["https://actual.example/product"])
    model = SequenceModel([missing, valid])
    provider = EvidenceBoundReportDraftProvider(model, FakeSearch())

    draft = provider.generate(fixture_task())

    assert draft.executive_summary_source_urls == ["https://actual.example/product"]
    assert len(model.prompts) == 3
    assert "引用修复任务" in model.prompts[1]
    assert "证据相关性审计" in model.prompts[2]


def test_report_provider_repairs_invalid_contract_once() -> None:
    model = SequenceModel(
        [
            {"report_title": "缺少字段"},
            report_payload(["https://actual.example/product"]),
        ]
    )
    provider = EvidenceBoundReportDraftProvider(model, FakeSearch())

    draft = provider.generate(fixture_task())

    assert len(draft.sections) == 5
    assert len(model.prompts) == 3
    assert "格式修复任务" in model.prompts[1]


def test_report_provider_repairs_citation_relevance_and_reaudits() -> None:
    valid = report_payload(["https://actual.example/product"])
    model = SequenceModel([valid, valid], audit_statuses=["partial", "supported"])
    provider = EvidenceBoundReportDraftProvider(model, FakeSearch())

    draft = provider.generate(fixture_task())

    assert draft.sections[0].source_urls == ["https://actual.example/product"]
    assert any("证据相关性修复任务" in prompt for prompt in model.prompts)
    assert sum("证据相关性审计" in prompt for prompt in model.prompts) == 2


def test_report_provider_normalizes_short_section_list_without_inventing_facts() -> None:
    payload = report_payload(["https://actual.example/product"])
    payload["sections"] = payload["sections"][:3]
    payload["sources"] = [{"unexpected": "model ledger is ignored"}]
    provider = EvidenceBoundReportDraftProvider(FakeModel(payload), FakeSearch())

    draft = provider.generate(fixture_task())

    assert len(draft.sections) == 5
    assert draft.sections[3].paragraphs == ["当前公开证据不足以形成可靠结论。"]
    assert [source.url for source in draft.sources] == ["https://actual.example/product"]


def test_identity_evidence_is_inherited_for_ecosystem_table() -> None:
    task = fixture_task()
    identity_url = "https://official.example/open-api"
    task.sources = [
        EvidenceSource(
            url=identity_url,
            title="开放平台官方文档",
            source_type="official_docs",
            supports=["ecosystem", "api"],
        )
    ]
    task.related_products = [
        RelatedProduct(
            name="衡法开放API",
            kind="api",
            relationship="向同企业产品提供接口能力。",
            verification_status="confirmed",
            source_urls=[identity_url],
        )
    ]
    provider = EvidenceBoundReportDraftProvider(
        FakeModel(report_payload(["https://actual.example/product"])),
        FakeSearch(),
    )

    draft = provider.generate(task)

    inherited = next(source for source in draft.sources if source.url == identity_url)
    assert inherited.inherited_from_identity is True
    assert inherited.evidence_level == "official_primary"
    assert inherited.accessed_at is not None


def test_unsupported_strong_conclusions_are_downgraded() -> None:
    payload = report_payload(["https://actual.example/product"])
    payload["sections"][0]["paragraphs"] = ["该产品已形成壁垒，并形成成熟的变现路径。"]
    payload["insights"][0]["analysis"] = "其行业领先地位和B端付费意愿已经得到证明。"
    provider = EvidenceBoundReportDraftProvider(FakeModel(payload), FakeSearch())

    draft = provider.generate(fixture_task())

    combined = " ".join(draft.sections[0].paragraphs + [draft.insights[0].analysis])
    assert "已形成壁垒" not in combined
    assert "成熟的变现路径" not in combined
    assert "行业领先" not in combined


def test_reviewed_logic_errors_and_unverified_regulatory_claims_are_downgraded() -> None:
    payload = report_payload(["https://actual.example/product"])
    payload["executive_summary"] = ["得理法律大模型通过国家网信办备案，表明具备变现能力。"]
    payload["sections"][0]["paragraphs"] = ["8年研发与成立时间存在矛盾，影响公信力。"]
    payload["insights"][0]["analysis"] = "行业势能存在虚高可能。"
    provider = EvidenceBoundReportDraftProvider(FakeModel(payload), FakeSearch())

    draft = provider.generate(fixture_task())
    combined = " ".join(
        draft.executive_summary
        + draft.sections[0].paragraphs
        + [draft.insights[0].analysis]
    )

    assert "通过国家网信办备案" not in combined
    assert "表明具备变现能力" not in combined
    assert "存在矛盾" not in combined
    assert "影响公信力" not in combined
    assert "虚高" not in combined
    assert "未取得监管公开名单" in combined


def test_timeline_conflict_is_recorded_without_calling_it_a_contradiction() -> None:
    class TimelineSearch(FakeSearch):
        def search(self, query: str, *, max_results: int = 5) -> list[SearchDocument]:
            return [
                SearchDocument(
                    query=query,
                    title="产品说明",
                    url="https://actual.example/product",
                    content="产品说明称由联合实验室历经8年研发推出。",
                    retrieval_status="readable",
                ),
                SearchDocument(
                    query=query,
                    title="科研机构合作记录",
                    url="https://www.example.ac.cn/cooperation",
                    content="2019年双方合作成立法律人工智能联合实验室。",
                    retrieval_status="readable",
                ),
            ]

    provider = EvidenceBoundReportDraftProvider(
        FakeModel(report_payload(["https://actual.example/product"])),
        TimelineSearch(),
    )

    draft = provider.generate(fixture_task())

    note = next(item for item in draft.limitations if "时间口径差异" in item)
    assert "历经8年研发" in note
    assert "2019年成立" in note
    assert "不能直接视为矛盾" in note


def test_platform_and_research_sources_are_not_labeled_official_primary() -> None:
    task = fixture_task()
    platform_url = "https://apps.apple.com/cn/app/example/id1"
    research_url = "https://www.example.ac.cn/news/legal-ai.html"
    task.sources = [
        EvidenceSource(
            url=platform_url,
            title="应用商店页面",
            source_type="official_product",
            supports=["identity"],
        ),
        EvidenceSource(
            url=research_url,
            title="科研机构合作信息",
            source_type="official_product",
            supports=["ecosystem"],
        ),
    ]
    payload = report_payload([platform_url])
    payload["sections"][0]["source_urls"] = [research_url]
    provider = EvidenceBoundReportDraftProvider(FakeModel(payload), FakeSearch())

    draft = provider.generate(task)
    levels = {source.url: source.evidence_level for source in draft.sources}

    assert levels[platform_url] == "platform_listing"
    assert levels[research_url] == "research_institution"


def test_word_builder_creates_editable_document(tmp_path) -> None:
    task = fixture_task()
    draft = FixtureReportDraftProvider().generate(task)
    output = tmp_path / "fixture.docx"

    build_report_docx(draft, task, output)
    document = Document(output)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert output.read_bytes().startswith(b"PK")
    assert "衡法智研标杆分析报告" in text
    assert "本报告仅用于自动化测试，不包含真实企业事实。  [1]" in text
    assert all(heading in text for heading in REPORT_SECTION_HEADINGS)
    assert "主要数据来源一览" in text
    assert "本节证据：[1]" in text
    assert "[1] 虚构自动化测试来源" in text

    source_heading_index = next(
        index
        for index, paragraph in enumerate(document.paragraphs)
        if paragraph.text == "主要数据来源一览"
    )
    preceding = document.paragraphs[source_heading_index - 1]
    page_breaks = preceding._p.findall(f".//{qn('w:br')}")
    assert any(node.get(qn("w:type")) == "page" for node in page_breaks)


def test_word_builder_downgrades_unmapped_confirmed_ecosystem_item(tmp_path) -> None:
    task = fixture_task()
    task.related_products = [
        RelatedProduct(
            name="无来源API",
            kind="api",
            relationship="模型声称存在关联。",
            verification_status="confirmed",
            source_urls=["https://not-in-ledger.example/api"],
        )
    ]
    draft = FixtureReportDraftProvider().generate(task)
    output = tmp_path / "unmapped.docx"

    build_report_docx(draft, task, output)
    document = Document(output)
    table_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)

    assert "仅发现公开线索" in table_text
    assert "已由公开来源确认" not in table_text
