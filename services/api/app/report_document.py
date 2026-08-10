from __future__ import annotations

from datetime import date
from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .models import ReportDraft, Task


BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(92, 102, 110)
BODY = RGBColor(30, 36, 40)
LIGHT_FILL = "F2F4F7"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120

SOURCE_LEVEL_LABELS = {
    "government_record": "政府或监管公开记录",
    "official_primary": "官方原始资料",
    "platform_listing": "应用商店或分发平台资料",
    "research_institution": "科研机构公开资料",
    "independent_secondary": "第三方公开资料",
    "other_public": "其他公开资料",
    "test_fixture": "虚构测试资料",
}
RELATED_STATUS_LABELS = {
    "confirmed": "已由公开来源确认",
    "lead_only": "仅发现公开线索",
    "not_publicly_verified": "未获公开证据确认",
}
RELATED_KIND_LABELS = {
    "product": "产品",
    "api": "API",
    "mcp": "MCP",
    "model_platform": "模型平台",
    "developer_platform": "开发者平台",
    "database": "数据库",
    "knowledge_base": "知识库",
    "agent": "Agent",
    "plugin": "插件",
    "connector": "连接器",
    "skill": "Skill",
    "other": "其他",
}


def _set_run_font(run, *, size: float, bold: bool = False, color=BODY) -> None:
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def _moderate_ecosystem_regulatory_claim(text: str, *, has_regulatory_source: bool) -> str:
    if has_regulatory_source or "备案" not in text:
        return text
    return "企业公开资料称其已完成相关备案，但本次未取得监管公开名单的直接对应记录。"


def _configure_styles(document: Document) -> None:
    tokens = {
        "Normal": (11, BODY, False, 0, 6, 1.10),
        "Heading 1": (16, BLUE, True, 16, 8, 1.0),
        "Heading 2": (13, BLUE, True, 12, 6, 1.0),
        "Heading 3": (12, DARK_BLUE, True, 8, 4, 1.0),
        "List Bullet": (11, BODY, False, 0, 8, 1.167),
        "List Number": (11, BODY, False, 0, 8, 1.167),
    }
    for name, (size, color, bold, before, after, spacing) in tokens.items():
        style = document.styles[name]
        style.font.name = "Arial"
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = spacing


def _set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_geometry(table, widths: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths, strict=True):
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)


def _shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def _add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    _set_run_font(run, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    run = paragraph.add_run(" 页")
    _set_run_font(run, size=9, color=MUTED)


def _safe_file_component(value: str) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", value).strip(" .")
    return cleaned[:60] or "法律AI产品"


def report_file_name(task: Task, generated_on: date) -> str:
    product = _safe_file_component(task.canonical_product_name or task.submitted_product_name)
    return f"{product}标杆分析报告_法律科技赛道.V0.4.P2.3.4测试稿.{generated_on:%Y%m%d}.docx"


def _source_reference_text(urls: list[str], source_numbers: dict[str, int]) -> str:
    refs: list[int] = []
    for url in urls:
        normalized = url.split("#", 1)[0].rstrip("/")
        number = source_numbers.get(normalized)
        if number is not None and number not in refs:
            refs.append(number)
    return "、".join(f"[{number}]" for number in refs)


def build_report_docx(
    draft: ReportDraft,
    task: Task,
    output_path: Path,
    *,
    generated_on: date | None = None,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    _configure_styles(document)
    section = document.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header.add_run("法律 AI 产品标杆分析 | P2.3.4 证据相关性审计测试稿")
    _set_run_font(run, size=9, color=MUTED)
    _add_page_number(section.footer.paragraphs[0])

    for _ in range(5):
        document.add_paragraph().paragraph_format.space_after = Pt(12)
    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = kicker.add_run("法律科技赛道 · 标杆分析")
    _set_run_font(run, size=11, bold=True, color=BLUE)
    kicker.paragraph_format.space_after = Pt(18)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(draft.report_title)
    _set_run_font(run, size=28, bold=True, color=DARK_BLUE)
    title.paragraph_format.space_after = Pt(10)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(draft.report_subtitle)
    _set_run_font(run, size=14, color=MUTED)
    subtitle.paragraph_format.space_after = Pt(48)
    metadata = document.add_paragraph()
    metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document_date = generated_on or date.today()
    meta_text = (
        f"产品：{task.canonical_product_name or task.submitted_product_name}\n"
        f"运营主体：{task.operator_name or '公开资料待补充'}\n"
        f"生成日期：{document_date:%Y年%m月%d日}"
    )
    for index, line in enumerate(meta_text.splitlines()):
        if index:
            metadata.add_run().add_break()
        run = metadata.add_run(line)
        _set_run_font(run, size=10.5, color=MUTED)
    document.add_page_break()

    source_numbers = {
        source.url.split("#", 1)[0].rstrip("/"): index
        for index, source in enumerate(draft.sources, start=1)
    }
    source_levels = {
        source.url.split("#", 1)[0].rstrip("/"): source.evidence_level
        for source in draft.sources
    }

    document.add_heading("报告摘要", level=1)
    for index, paragraph in enumerate(draft.executive_summary):
        summary_p = document.add_paragraph()
        summary_p.add_run(paragraph)
        item_urls = (
            draft.executive_summary_item_source_urls[index]
            if index < len(draft.executive_summary_item_source_urls)
            else draft.executive_summary_source_urls
        )
        summary_references = _source_reference_text(item_urls, source_numbers)
        if summary_references:
            run = summary_p.add_run(f"  {summary_references}")
            _set_run_font(run, size=8.5, color=MUTED)

    scope_note = document.add_paragraph()
    scope_run = scope_note.add_run(
        "证据口径：政府或监管记录用于核验公开登记事实；企业官网、官方文档和应用说明仅证明其公开披露；分析判断不等同于已核实事实。"
    )
    _set_run_font(scope_run, size=8.5, color=MUTED)

    if task.related_products:
        document.add_heading("产品生态速览", level=2)
        table = document.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        headers = ["产品或入口", "类型", "核验状态", "与核心产品的关系及证据"]
        for cell, text in zip(table.rows[0].cells, headers, strict=True):
            _shade_cell(cell, LIGHT_FILL)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            _set_run_font(run, size=9.5, bold=True, color=DARK_BLUE)
        _set_repeat_table_header(table.rows[0])
        for item in task.related_products:
            cells = table.add_row().cells
            references = _source_reference_text(item.source_urls, source_numbers)
            effective_status = item.verification_status
            if effective_status == "confirmed" and not references:
                effective_status = "lead_only"
            has_regulatory_source = any(
                source_levels.get(url.split("#", 1)[0].rstrip("/")) == "government_record"
                for url in item.source_urls
            )
            relation = _moderate_ecosystem_regulatory_claim(
                item.relationship,
                has_regulatory_source=has_regulatory_source,
            )
            if references:
                relation = f"{relation} {references}"
            values = [
                item.name,
                RELATED_KIND_LABELS.get(item.kind, item.kind),
                RELATED_STATUS_LABELS.get(effective_status, effective_status),
                relation,
            ]
            for index, (cell, text) in enumerate(zip(cells, values, strict=True)):
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if index in {1, 2} else WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(text)
                _set_run_font(run, size=9.5)
        _set_table_geometry(table, [1800, 1200, 2040, 4320])

    for section_draft in draft.sections:
        document.add_heading(section_draft.heading, level=1)
        for paragraph in section_draft.paragraphs:
            document.add_paragraph(paragraph)
        for point in section_draft.key_points:
            document.add_paragraph(point, style="List Bullet")
        if section_draft.source_urls:
            source_p = document.add_paragraph()
            source_p.paragraph_format.space_before = Pt(4)
            source_p.paragraph_format.space_after = Pt(4)
            references = _source_reference_text(section_draft.source_urls, source_numbers)
            run = source_p.add_run("本节证据：" + (references or "未形成有效来源映射"))
            _set_run_font(run, size=8.5, color=MUTED)

    document.add_heading("六、洞察总结", level=1)
    for index, insight in enumerate(draft.insights, start=1):
        p = document.add_paragraph()
        number_run = p.add_run(f"{index}.  ")
        _set_run_font(number_run, size=11)
        title_run = p.add_run(f"{insight.title}：")
        _set_run_font(title_run, size=11, bold=True, color=DARK_BLUE)
        body_run = p.add_run(insight.analysis)
        _set_run_font(body_run, size=11)
        references = _source_reference_text(insight.source_urls, source_numbers)
        if references:
            source_run = p.add_run(f" {references}")
            _set_run_font(source_run, size=8.5, color=MUTED)

    if draft.limitations:
        document.add_heading("研究边界", level=1)
        for item in draft.limitations:
            document.add_paragraph(item, style="List Bullet")

    document.add_page_break()
    document.add_heading("主要数据来源一览", level=1)
    for index, source in enumerate(draft.sources, start=1):
        p = document.add_paragraph()
        level = SOURCE_LEVEL_LABELS.get(source.evidence_level, "其他公开资料")
        access = f"；访问日期：{source.accessed_at}" if source.accessed_at else ""
        inherited = "；身份核验阶段继承" if source.inherited_from_identity else ""
        run = p.add_run(f"[{index}] {source.title}（{level}{access}{inherited}）\n{source.url}")
        _set_run_font(run, size=9.5)

    document.core_properties.title = draft.report_title
    document.core_properties.subject = "法律AI产品标杆分析P2.3.4证据相关性审计测试稿"
    document.core_properties.author = "法律AI分析Agent"
    document.core_properties.comments = "本地生成；不含数据库写入或训练。"
    document.save(output_path)
